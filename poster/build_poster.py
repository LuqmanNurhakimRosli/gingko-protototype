# Ginkgo poster for URBIEE 2026: rewrite the poster to match the current system.
#   python poster/build_poster.py
# Starts from "GINKGO poster for URBIEE26 (before update).pptx" (a copy of the original, made on the first run),
# keeps the layout, fonts, colours, header logos and author block, and replaces the Ginkgo logo, the text,
# the two charts (now native, from the system's own numbers) and adds a "System in action" band.
#   needs: poster/data/drill_series.json (node poster/capture_drill_series.mjs), poster/assets/* (node poster/make_assets.mjs),
#          public/data/stats.json, and the captured app frames in demo/pitch/work/capture/ (npm run pitch:capture).

import copy
import io
import json
import shutil
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.shapes.group import GroupShape
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
POSTER = ROOT / 'poster'
OUT = POSTER / 'GINKGO poster for URBIEE26.pptx'
BACKUP = POSTER / 'GINKGO poster for URBIEE26 (before update).pptx'
CAPTURE = ROOT / 'demo' / 'pitch' / 'work' / 'capture'

TEAL = RGBColor(0x07, 0x5B, 0x50)
INK = RGBColor(0x22, 0x22, 0x22)
RED, AMBER, BLUE, GREEN, PURPLE = (RGBColor.from_string(c) for c in ('D93025', 'EDA100', '3884EC', '08A045', 'AA00AA'))
TINT = RGBColor.from_string('E4F0EC')


def pt(v):
    return Emu(int(round(v * 12700)))


# ── content ────────────────────────────────────────────────────────────────────────────────────────
TITLE = 'GINKGO'
SUBTITLE = 'FLOOD COMMAND CENTER: FROM RISK MAP TO EVACUATION PLAN FOR SELANGOR'
ABSTRACT = [
    "When floods come, people hear that a warning is in force, but not when to leave or where to go, and nobody can see which "
    "shelters will fill up. Ginkgo is a web-based flood command center for Selangor's nine districts that closes this gap.",
    "Planning mode maps where river flooding is most likely, from satellite terrain, rivers, 38 years of surface-water records "
    "and census data on a grid of about 92 m cells. Flood drill mode then replays a storm hour by hour: it warns each neighbourhood on Japan's "
    "five-level scale, finds a walking or driving route along roads that stay dry to a centre that still has space, and says "
    "when to leave. A 3D model of Taman Sri Muda shows the water on the street, and an AI assistant answers in English or Malay "
    "with every figure checked against the data.",
    "It runs in a browser on open data (Copernicus, OpenStreetMap, JRC, ESA, Meta, DOSM), so no GIS software is needed. "
    "The drill is a simulation for rehearsal, not a forecast.",
]
OBJECTIVES = [
    "Map where river flooding is most likely in Selangor under Normal, Heavy and Extreme rain, in a form anyone can read",
    "Rehearse a flood hour by hour, with alerts on Japan's five-level scale and a one-hour look-ahead for each neighbourhood",
    "Guide each household along a dry route to an evacuation centre that still has space, and say when to leave",
    "Show the water on the street in 3D, and let anyone ask in English or Malay with every number checked against the data",
]
STATS = [
    ('24.1%', "of Selangor's land has High flood susceptibility"),
    ('~1.93M', 'people live on land at High risk (estimate)'),
    ('1,274', 'candidate evacuation centres, each with a capacity'),
]
USEFUL = [  # (text, lines at this width)
    ("Turns open satellite and census data into a risk map, a drill and routes that anyone can use", 3),
    ("Gives each household a dry route, a centre with space and a leave-by time", 2),
    ("Open data and a browser app keep costs low: an AI answer is a few US cents, repeats are free", 3),
    ("Next: a council pilot on live river gauges, then Malaysia and ASEAN", 2),
]
CONCLUSION = (
    "Ginkgo turns a flood map into a plan: when each neighbourhood must leave, where there is space, and how to get there on "
    "dry roads. Checked against 7 reported December 2021 flood sites (4 flagged High, 2 of 2 hilly controls Low). "
    "A simulation for rehearsal, not a forecast."
)
SHOTS = [
    ('1  MAP', 'Where river flooding is most likely'),
    ('2  DRILL', 'Alert, dry route, centre with space'),
    ('3  3D', 'The water on the street'),
]


# ── helpers ────────────────────────────────────────────────────────────────────────────────────────
def walk(shapes):
    for s in shapes:
        yield s
        if isinstance(s, GroupShape):
            yield from walk(s.shapes)


def by_name(slide, name):
    for s in walk(slide.shapes):
        if s.name == name:
            return s
    raise KeyError(name)


def remove(shape):
    shape._element.getparent().remove(shape._element)


def place(shape, x=None, y=None, w=None, h=None):
    if x is not None: shape.left = pt(x)
    if y is not None: shape.top = pt(y)
    if w is not None: shape.width = pt(w)
    if h is not None: shape.height = pt(h)


def set_text(shape, text, size=None, like=None, gap=0):
    """Replace a text box's paragraphs, keeping the first paragraph's formatting (bullets, spacing, font).
    like: copy that formatting from another text box instead; gap: points before each paragraph after the first."""
    body = shape.text_frame._txBody
    paras = body.findall(qn('a:p'))
    tmpl = copy.deepcopy((like.text_frame._txBody.findall(qn('a:p')) if like is not None else paras)[0])
    runs = tmpl.findall(qn('a:r'))
    rpr = copy.deepcopy(runs[0].find(qn('a:rPr')))
    if size:
        rpr.set('sz', str(int(size * 100)))
    for r in runs:
        tmpl.remove(r)
    for e in tmpl.findall(qn('a:endParaRPr')):
        tmpl.remove(e)
    for p in paras:
        body.remove(p)
    for i, t in enumerate([text] if isinstance(text, str) else text):
        p = copy.deepcopy(tmpl)
        if gap and i:
            sb = etree.Element(qn('a:spcBef'))
            etree.SubElement(sb, qn('a:spcPts')).set('val', str(int(gap * 100)))
            p.find(qn('a:pPr')).find(qn('a:lnSpc')).addnext(sb)
        r = etree.SubElement(p, qn('a:r'))
        r.append(copy.deepcopy(rpr))
        etree.SubElement(r, qn('a:t')).text = t
        body.append(p)


def add_text(slide, x, y, w, h, text, font='Montserrat', size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, name=None, spc=None, line=None):
    tb = slide.shapes.add_textbox(pt(x), pt(y), pt(w), pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if line:
        p.line_spacing = Pt(line)
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold = font, Pt(size), bold
    f.color.rgb = color
    if spc is not None:
        r._r.get_or_add_rPr().set('spc', str(spc))
    if name:
        tb.name = name
    return tb


def alt(shape, text):
    shape._element.xpath('.//p:cNvPr')[0].set('descr', text)


def crop_save(src, box, dst, width):
    im = Image.open(src).convert('RGB').crop(box)
    im = im.resize((width, round(im.height * width / im.width)), Image.LANCZOS)
    im.save(dst, quality=90)
    return dst


# ── build ──────────────────────────────────────────────────────────────────────────────────────────
def main():
    if not BACKUP.exists():
        shutil.copy2(OUT, BACKUP)
    prs = Presentation(str(BACKUP))
    sl = prs.slides[0]

    stats = json.loads((ROOT / 'public' / 'data' / 'stats.json').read_text(encoding='utf-8'))
    series = json.loads((POSTER / 'data' / 'drill_series.json').read_text(encoding='utf-8'))
    nrm = stats['state']['scenarios']['normal']
    assert (nrm['high_pct'], round(nrm['pop_high'] / 1e6, 2)) == (24.1, 1.93), 'poster figures out of date'

    # header: the Ginkgo logo (the old lockup picture is used twice: here and under the objective)
    lock = Image.open(POSTER / 'assets' / 'ginkgo-lockup.png')
    l, t, r, b = lock.getchannel('A').getbbox()
    lock = lock.crop((max(0, l - 24), max(0, t - 24), min(lock.width, r + 24), min(lock.height, b + 24)))
    buf = io.BytesIO(); lock.save(buf, 'PNG')
    sl.part.related_part('rId17')._blob = buf.getvalue()
    logo = by_name(sl, 'Freeform 57')
    lh = 118  # the slot between the URBIEE and UiTM logos is ~500 pt wide
    place(logo, x=972 - lh * lock.width / lock.height / 2, y=101, w=lh * lock.width / lock.height, h=lh)
    alt(by_name(sl, 'Freeform 57'), 'Ginkgo logo: a ginkgo leaf holding a city skyline above floodwater, with the words GINKGO, flood command center')
    remove(by_name(sl, 'Freeform 84'))

    title = by_name(sl, 'TextBox 58')
    set_text(title, TITLE)
    set_text(by_name(sl, 'TextBox 76'), SUBTITLE)

    # abstract and objective
    set_text(by_name(sl, 'TextBox 59'), ABSTRACT, gap=12)
    first = by_name(sl, 'TextBox 60')
    for name, text in zip(('TextBox 60', 'TextBox 79', 'TextBox 80', 'TextBox 81'), OBJECTIVES):
        set_text(by_name(sl, name), text, like=first)
        place(by_name(sl, name), x=870, w=670)  # the original's last two bullets sat 4 pt left of the first two

    # key figures under the objective (where the second logo was)
    x0, y0, gap = 866, 1088, 14
    tw, th = (714 - 2 * gap) / 3, 186
    for i, (big, small) in enumerate(STATS):
        x = x0 + i * (tw + gap)
        tile = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pt(x), pt(y0), pt(tw), pt(th))
        tile.adjustments[0] = 0.07
        tile.fill.solid(); tile.fill.fore_color.rgb = TINT
        tile.line.fill.background()
        tile.shadow.inherit = False
        tile.name = f'Key figure {i + 1}'
        add_text(sl, x + 14, y0 + 16, tw - 28, 70, big, font='Anton', size=48, color=TEAL, name=f'Key figure {i + 1} value', spc=60)
        add_text(sl, x + 14, y0 + 90, tw - 28, 86, small, size=16, color=INK, name=f'Key figure {i + 1} label', line=21)

    # usefulness and commercialisation: left bullets, right two charts
    panel_h = 436
    grp15, grp21 = by_name(sl, 'Group 15'), by_name(sl, 'Group 21')
    place(grp15, h=panel_h + 60)
    place(grp21, h=panel_h)
    for name in ('Picture 24', 'Picture 25', 'Group 26', 'Group 29', 'Group 32', 'TextBox 67', 'TextBox 68', 'TextBox 69'):
        remove(by_name(sl, name))
    boxes = [by_name(sl, n) for n in ('TextBox 71', 'TextBox 72', 'TextBox 73')]
    clone = copy.deepcopy(boxes[1]._element)  # a fourth bullet, same formatting
    clone.xpath('.//p:cNvPr')[0].set('id', str(max(int(i) for i in sl.shapes._spTree.xpath('.//p:cNvPr/@id')) + 1))
    clone.xpath('.//p:cNvPr')[0].set('name', 'TextBox 85')
    boxes[1]._element.addnext(clone)
    boxes = [by_name(sl, n) for n in ('TextBox 71', 'TextBox 72', 'TextBox 85', 'TextBox 73')]
    y = 1394
    for b, (text, lines) in zip(boxes, USEFUL):
        set_text(b, text)
        place(b, y=y, w=500, h=lines * 31.3)
        y += lines * 31.3 + 26

    set_text(by_name(sl, 'TextBox 65'), 'RISK BY DISTRICT')
    set_text(by_name(sl, 'TextBox 66'), 'THE EVACUATION GAP')
    set_text(by_name(sl, 'TextBox 82'), 'High and Moderate share of land')
    set_text(by_name(sl, 'TextBox 83'), 'Klang storm drill, people (thousands)')
    for name, x in (('TextBox 65', 696), ('TextBox 66', 1112)):
        place(by_name(sl, name), x=x, y=1394)
    place(by_name(sl, 'TextBox 82'), x=696, y=1438, w=403)
    place(by_name(sl, 'TextBox 83'), x=1112, y=1438, w=467)
    for name in ('TextBox 82', 'TextBox 83'):
        for p in by_name(sl, name).text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
    place(by_name(sl, 'TextBox 65'), w=403)
    place(by_name(sl, 'TextBox 66'), w=467)

    chart_y, chart_h = 1470, 316
    # chart 1: High and Moderate share of land by district, Normal rain
    rows = sorted(((d['name'], d['scenarios']['normal']['high_pct'], d['scenarios']['normal']['moderate_pct']) for d in stats['daerah']),
                  key=lambda r: r[1])  # ascending: a bar chart draws the first category at the bottom
    cd = CategoryChartData()
    cd.categories = [r[0] for r in rows]
    cd.add_series('High', [r[1] for r in rows])
    cd.add_series('Moderate', [r[2] for r in rows])
    gf = sl.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, pt(680), pt(chart_y), pt(430), pt(chart_h), cd)
    gf.name = 'Chart: flood risk by district'
    alt(gf, 'Stacked bar chart of the share of land with High and Moderate flood susceptibility in each of Selangor\'s nine districts, '
            'Normal rain. Klang is highest: 59.8% High and 32.1% Moderate.')
    c = gf.chart
    c.font.name, c.font.size = 'Montserrat', Pt(13)
    c.font.color.rgb = INK
    c.has_title = False
    c.has_legend = True
    c.legend.position, c.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    c.legend.font.size = Pt(14)
    plot = c.plots[0]
    plot.gap_width, plot.overlap = 35, 100
    for s, col in zip(plot.series, (RED, AMBER)):
        s.format.fill.solid(); s.format.fill.fore_color.rgb = col
    hi = plot.series[0].data_labels
    hi.show_value, hi.number_format, hi.number_format_is_linked, hi.position = True, '0"%"', False, XL_LABEL_POSITION.CENTER
    hi.font.size, hi.font.bold = Pt(12), True
    hi.font.color.rgb = RGBColor(255, 255, 255)
    va, ca = c.value_axis, c.category_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 100, 25
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '0"%"', False
    va.tick_labels.font.size = Pt(12)
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
    va.format.line.fill.background()
    ca.tick_labels.font.size = Pt(13)
    ca.format.line.color.rgb = RGBColor(0x99, 0x99, 0x99)

    # chart 2: the evacuation gap over the drill (the app's own headline tiles, every half hour)
    rs = series['rows']
    cd = CategoryChartData()
    cd.categories = [f'{r["minute"] / 60:g}' for r in rs]
    cd.add_series('Flooding within 1 h', [round(r['within1h'] / 1000) for r in rs])
    cd.add_series('In centres', [round(r['inCentres'] / 1000) for r in rs])
    cd.add_series('No centre in reach', [round(r['noCentre'] / 1000) for r in rs])
    gf = sl.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, pt(1118), pt(chart_y), pt(455), pt(chart_h), cd)
    gf.name = 'Chart: the evacuation gap'
    alt(gf, 'Line chart over the first five hours of the Klang Valley storm drill, in thousands of people: the number whose cells flood within '
            'one hour peaks near 633 thousand at 3 hours; centres hold up to 141 thousand; 206 thousand have no centre within reach.')
    c = gf.chart
    c.font.name, c.font.size = 'Montserrat', Pt(13)
    c.font.color.rgb = INK
    c.has_title = False
    c.has_legend = True
    c.legend.position, c.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    c.legend.font.size = Pt(13)
    for s, col in zip(c.plots[0].series, (PURPLE, GREEN, RED)):
        s.smooth = False
        s.format.line.color.rgb = col
        s.format.line.width = Pt(3.5)
        s.marker.style, s.marker.size = XL_MARKER_STYLE.CIRCLE, 7
        s.marker.format.fill.solid(); s.marker.format.fill.fore_color.rgb = col
        s.marker.format.line.color.rgb = col
    va, ca = c.value_axis, c.category_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 700, 100
    va.tick_labels.font.size = Pt(12)
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
    va.format.line.fill.background()
    ca.tick_labels.font.size = Pt(12)
    ca.has_title = True
    ca.axis_title.text_frame.text = 'hours since the rain started'
    ca.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    ca.axis_title.text_frame.paragraphs[0].runs[0].font.bold = False
    ca.format.line.color.rgb = RGBColor(0x99, 0x99, 0x99)

    add_text(sl, 696, 1790, 403, 20, 'Source: Ginkgo analysis of open data', size=14, color=RGBColor(0x55, 0x55, 0x55),
             align=PP_ALIGN.CENTER, name='Source: risk by district')
    add_text(sl, 1112, 1790, 467, 20, 'Source: Ginkgo flood drill, Extreme rain (estimate)', size=14, color=RGBColor(0x55, 0x55, 0x55),
             align=PP_ALIGN.CENTER, name='Source: evacuation gap')

    # system in action: three real screenshots
    band_y, band_h = 1835, 270
    box = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pt(115), pt(band_y), pt(1465), pt(band_h))
    box.adjustments[0] = 0.045
    box.fill.background(); box.line.color.rgb = TEAL; box.line.width = Pt(2.25)
    box.shadow.inherit = False
    box.name = 'System in action: frame'
    badge = sl.shapes.add_shape(MSO_SHAPE.OVAL, pt(115), pt(band_y + 2), pt(73), pt(71))
    badge.fill.solid(); badge.fill.fore_color.rgb = TEAL; badge.line.fill.background(); badge.shadow.inherit = False
    badge.name = 'System in action: badge'
    icon = sl.shapes.add_picture(str(POSTER / 'assets' / 'icon-monitor.png'), pt(133), pt(band_y + 19), pt(37), pt(37))
    alt(icon, 'Screen icon')
    add_text(sl, 208, band_y + 14, 700, 46, 'SYSTEM IN ACTION', font='Anton', size=34.93, color=TEAL, spc=216, name='System in action: title',
             line=48.9)
    rule = sl.shapes.add_connector(1, pt(136), pt(band_y + 72), pt(1579), pt(band_y + 72))
    rule.line.color.rgb, rule.line.width = TEAL, Pt(3)
    rule.name = 'System in action: rule'

    asp = 2.7
    crops = [
        (CAPTURE / 's1' / 'f_0000.jpg', 1.5, (320, 470), 1180),
        (CAPTURE / 's3' / 'f_0149.jpg', 1.5, (320, 560), 1180),
        (CAPTURE / 's5' / 'f_0215.jpg', 1.0, (0, 250), 1920),
    ]
    iw, ih, pad = 459, 170, 22
    descr = ['Screenshot of the planning map of Selangor, red where river flooding is most likely',
             'Screenshot of the flood drill: a blue evacuation route on the map, around flooded neighbourhoods, to a centre with space',
             'Screenshot of the 3D model of Taman Sri Muda with floodwater between the houses']
    for i, ((src, k, (cx, cy), cw), (tag, cap), d) in enumerate(zip(crops, SHOTS, descr)):
        ch = round(cw / asp)
        dst = crop_save(src, (round(cx * k), round(cy * k), round((cx + cw) * k), round((cy + ch) * k)),
                        POSTER / 'assets' / f'shot{i + 1}.jpg', 1380)
        x = 115 + pad + i * (iw + pad)
        pic = sl.shapes.add_picture(str(dst), pt(x), pt(band_y + 72 + 16), pt(iw), pt(ih))
        pic.line.color.rgb, pic.line.width = RGBColor(0xB8, 0xC4, 0xC1), Pt(1)
        pic.name = f'Screenshot {i + 1}'
        alt(pic, d)
        pill = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pt(x + 10), pt(band_y + 72 + 16 + ih - 44), pt(iw - 20), pt(34))
        pill.adjustments[0] = 0.3
        pill.fill.solid(); pill.fill.fore_color.rgb = TEAL; pill.line.fill.background(); pill.shadow.inherit = False
        pill.name = f'Screenshot {i + 1} caption'
        tf = pill.text_frame
        tf.margin_left = tf.margin_right = pt(12); tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        r1 = p.add_run(); r1.text = tag + '   '
        r1.font.name, r1.font.size, r1.font.bold = 'Montserrat Bold', Pt(16), True
        r1.font.color.rgb = RGBColor(0xB7, 0xE8, 0xC5)
        r2 = p.add_run(); r2.text = cap
        r2.font.name, r2.font.size, r2.font.bold = 'Montserrat', Pt(16), False
        r2.font.color.rgb = RGBColor(255, 255, 255)

    # conclusion, moved up and made shorter to make room
    top = band_y + band_h + 20
    place(by_name(sl, 'Group 35'), y=top, h=2351 - top)
    place(by_name(sl, 'Group 46'), y=top + 2)
    place(by_name(sl, 'Freeform 49'), y=top + 12)
    place(by_name(sl, 'TextBox 75'), y=top + 14)
    place(by_name(sl, 'AutoShape 50'), y=top + 72)
    set_text(by_name(sl, 'TextBox 64'), CONCLUSION)
    place(by_name(sl, 'TextBox 64'), y=top + 72 + 17, h=3 * 31.3)

    prs.save(str(OUT))
    print('wrote', OUT)


main()
